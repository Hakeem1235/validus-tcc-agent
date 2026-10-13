"""S1 intake: decode attachments -> text (PDF/DOCX parse, images via vision LLM) -> per-document facts."""
from __future__ import annotations

import base64
import io
import re

from app import llm
from app.normalise import iso, last4, parse_date
from app.schemas import Attachment, AttachmentInfo

IMAGE_MIMES = {"image/jpeg", "image/jpg", "image/png", "image/gif", "image/webp", "image/tiff", "image/bmp"}

VISION_SCHEMA = {
    "type": "object",
    "properties": {
        "transcription": {"type": "string", "description": "verbatim text of the document, top to bottom"},
        "document_kind": {"type": "string"},
        "ticked_boxes": {"type": "array", "items": {"type": "string"}, "description": "labels of checkboxes that are ticked"},
        "unticked_boxes": {"type": "array", "items": {"type": "string"}},
        "account_number": {"type": ["string", "null"]},
        "account_holder_names": {"type": "array", "items": {"type": "string"}},
        "ifsc": {"type": ["string", "null"]},
        "folio_number": {"type": ["string", "null"]},
        "signed": {"type": ["boolean", "null"]},
        "date": {"type": ["string", "null"]},
    },
    "required": ["transcription", "document_kind", "ticked_boxes", "account_holder_names"],
}

VISION_PROMPT = (
    "You are transcribing a scanned document attached to an email sent to a mutual-fund registrar. "
    "Transcribe all printed and handwritten text exactly. List checkbox labels that are ticked and those that are not. "
    "Extract the account number, account holder name(s) printed on it, IFSC, folio number and date if present. "
    "Never guess: use null when a value is not visible."
)

# doc-type rules: (type, pattern on text or file name)
DOC_RULES: list[tuple[str, str]] = [
    ("TPP_DECLARATION", r"third[\s_-]*party[\s_-]*payment[\s_-]*declaration|tpp[\s_-]*declaration"),
    ("COB_FORM", r"change[\s_-]*of[\s_-]*bank|cob[\s_-]*request|bank[\s_-]*mandate[\s_-]*change"),
    ("BANK_STATEMENT", r"statement[\s_-]*of[\s_-]*account|bank[\s_-]*statement|account[\s_-]*statement"),
    ("TRANSFER_INSTRUCTION", r"funds?[\s_-]*transfer|transaction[\s_-]*acknowledg|neft[\s_-]*advice|rtgs[\s_-]*advice|"
                             r"imps[\s_-]*advice|transfer[\s_-]*instruction|debit[\s_-]*advice"),
    ("CANCELLED_CHEQUE", r"cancel+ed[\s_-]*cheque|\bcancel+ed\b.*\ba/c\b|cheque.*cancel+ed"),
    ("PASSBOOK", r"pass[\s_-]*book"),
    ("BANK_LETTER", r"bank[\s_-]*letter|to[\s_-]*whomsoever"),
    ("BANKER_CERTIFICATE", r"banker'?s[\s_-]*certificate"),
    ("EMPLOYER_DECLARATION", r"payroll[\s_-]*deduction|employee[\s_-]*investment[\s_-]*programme|employer[\s_-]*declaration|payroll[\s_-]*letter"),
    ("EMPLOYEE_LIST", r"^\s*(list|schedule)\s+of\s+employees"),
    ("KYC_ACK_INVESTOR", r"kyc[\s_-]*acknowledg\w*[\s_-]*(of[\s_-]*)?(the[\s_-]*)?(investor|employee)"),
    ("KYC_ACK_THIRD_PARTY", r"kyc[\s_-]*acknowledg\w*[\s_-]*(of[\s_-]*)?(the[\s_-]*)?(employer|third[\s_-]*party|company|custodian|corporate)"),
    ("EXECUTION_ONLY_DECLARATION", r"execution[\s_-]*only"),
    ("CHANGE_OF_STATUS_FORM", r"change[\s_-]*of[\s_-]*(tax[\s_-]*)?status|minor[\s_-]*to[\s_-]*major"),
    ("REDEMPTION_FORM", r"redemption[\s_-]*(request|form|slip)"),
    ("CUSTODIAN_SEBI_REG", r"custodian.*sebi[\s_-]*registration"),
    ("DISTRIBUTOR_LIST", r"list[\s_-]*of[\s_-]*distributors"),
]


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return "\n".join((p.extract_text() or "") for p in pdf.pages)


def _docx_text(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            lines.append(" | ".join(c.text for c in r.cells))
    return "\n".join(lines)


def _vision(data: bytes, mime: str, info: AttachmentInfo) -> None:
    out = llm.generate_json(VISION_PROMPT + f"\nFile name: {info.file_name}", purpose="vision_transcription",
                            json_schema=VISION_SCHEMA, images=[(data, mime)], max_output_tokens=2048)
    info.method = "vision"
    info.text = out.get("transcription") or ""
    info.ticked = out.get("ticked_boxes") or []
    info.flags["unticked"] = out.get("unticked_boxes") or []
    info.flags["document_kind"] = out.get("document_kind")
    if out.get("account_number"):
        info.account_last4 = last4(out["account_number"])
    info.account_holders = [h for h in out.get("account_holder_names") or [] if h]
    info.ifsc = out.get("ifsc") or info.ifsc
    info.folio_raw = out.get("folio_number") or info.folio_raw
    if out.get("date"):
        info.doc_date = iso(parse_date(out["date"])) or info.doc_date


def classify(info: AttachmentInfo) -> None:
    """Assign document types from content first, file name second."""
    text = info.text or ""
    head = "\n".join(text.splitlines()[:6])
    name = info.file_name.replace("_", " ")
    types: list[str] = []
    for t, pat in DOC_RULES:
        if t == "EMPLOYEE_LIST":
            if re.search(pat, text, re.I | re.M) and not re.search(r"enclosure\s*:\s*list of employees", text, re.I):
                types.append(t)
            continue
        if t.startswith("KYC_ACK"):
            # a KYC acknowledgement is a document in its own right, not a checkbox label on another form
            if re.search(pat, head, re.I) or re.search(pat, name, re.I):
                types.append(t)
            continue
        if t in ("EXECUTION_ONLY_DECLARATION", "CHANGE_OF_STATUS_FORM", "CUSTODIAN_SEBI_REG", "DISTRIBUTOR_LIST"):
            if re.search(pat, head, re.I) or re.search(pat, name, re.I):
                types.append(t)
            continue
        if re.search(pat, head, re.I) or re.search(pat, name, re.I) or (t in ("BANK_STATEMENT", "TRANSFER_INSTRUCTION", "COB_FORM", "EMPLOYER_DECLARATION") and re.search(pat, text[:400], re.I)):
            types.append(t)
    if "CANCELLED_CHEQUE" in types and "COB_FORM" in types and not re.search(r"cheque", name, re.I) and "COB_FORM" in types:
        # a COB form listing "cancelled cheque" as an enclosure is not the cheque itself
        if re.search(r"enclosures?\s*:", text, re.I) and not re.search(r"cheque", head, re.I):
            types.remove("CANCELLED_CHEQUE")
    if not types:
        types = ["OTHER"]
    info.doc_types = list(dict.fromkeys(types))


def parse_fields(info: AttachmentInfo) -> None:
    """Deterministic field extraction from attachment text."""
    t = info.text or ""
    if not t:
        return
    acct_pats = [r"(?:debit|remitter|account|a/c|acct)\s*(?:no\.?|number|#)?\s*[:\-]?\s*([X\*\d][X\*\d\s]{5,}\d)",
                 r"A/c\s*No\.?\s*([\dX\*]{6,})"]
    if not info.account_last4:
        for p in acct_pats:
            m = re.search(p, t, re.I)
            if m:
                info.account_last4 = last4(m.group(1))
                break
    holders: list[str] = []
    for p in [r"account\s*holder\s*(?:name)?\s*[:\-]\s*([A-Z][A-Za-z .]+)", r"remitter\s*name\s*[:\-]\s*([A-Z][A-Za-z .]+)",
              r"joint\s*holder\(?s?\)?\s*[:\-]\s*([A-Z][A-Za-z .;,]+)"]:
        for m in re.finditer(p, t, re.I):
            for h in re.split(r"[;,]", m.group(1)):
                h = h.strip(" .")
                if h and h not in ("—", "-") and len(h) > 2:
                    holders.append(h)
    if "CANCELLED_CHEQUE" in info.doc_types and not holders:
        m = re.search(r"A/c\s*No\.?\s*[\dX\*]+\s*\n\s*([A-Z][A-Z .]{3,})", t)
        if m:
            holders.append(m.group(1).strip())
    if not info.account_holders:
        info.account_holders = list(dict.fromkeys(holders))
    m = re.search(r"\b([A-Z]{4}0[A-Z0-9]{6})\b", t)
    if m and not info.ifsc:
        info.ifsc = m.group(1)
    m = re.search(r"\bUTR\s*[:\-]?\s*([A-Z0-9]{10,22})\b", t) or re.search(r"(?:NEFT|RTGS|IMPS)[-/ ]([A-Z]{4,5}[A-Z0-9]{8,18})", t)
    if m:
        info.utr = m.group(1)
    m = re.search(r"(?:amount|inr|rs\.?)\s*[:\-]?\s*(?:inr|rs\.?)?\s*([\d,]+(?:\.\d{1,2})?)", t, re.I)
    if m:
        try:
            info.amount_inr = float(m.group(1).replace(",", ""))
        except ValueError:
            pass
    m = re.search(r"\bfolio(?:\s*no\.?|\s*/\s*application\s*no\.?)?\s*[:\-]?\s*(\d{8}\s*[/\-]?\s*\d{2})\b", t, re.I)
    if m:
        info.folio_raw = re.sub(r"\s", "", m.group(1))
    m = re.search(r"\bdate\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{4}|\d{4}-\d{2}-\d{2}|\d{1,2}\s+\w+\s+\d{4})", t, re.I)
    if m and not info.doc_date:
        info.doc_date = iso(parse_date(m.group(1)))
    if "COB_FORM" in info.doc_types:
        info.flags["existing_proof_missing"] = bool(re.search(r"(old|existing)\s*bank\s*proof\s*:\s*(not|nil|none|n/?a)", t, re.I))
        info.flags["in_person"] = bool(re.search(r"in[\s-]*person\s*verif\w*\s*(done|completed|yes)", t, re.I))
        m = re.search(r"new\s*bank.*?account\s*number\s*:\s*([X\*\d]+)", t, re.I | re.S)
        if m:
            info.flags["new_account_last4"] = last4(m.group(1))
        m = re.search(r"existing\s*bank.*?(?:account|a/c)\s*:?\s*[A-Za-z ]*?([X\*\d]{6,})", t, re.I | re.S)
        if m:
            info.flags["existing_account_last4"] = last4(m.group(1))
        info.account_last4 = info.flags.get("new_account_last4") or info.account_last4
        m = re.search(r"received\s*at.*?(\d{1,2}/\d{1,2}/\d{4})", t, re.I)
        if m:
            info.doc_date = iso(parse_date(m.group(1)))


def intake(attachments: list[Attachment]) -> list[AttachmentInfo]:
    out = []
    for a in attachments:
        mime = (a.content_type or "").lower()
        name = a.file_name or "attachment"
        info = AttachmentInfo(file_name=name, content_type=mime)
        try:
            data = base64.b64decode(a.content_base64 or "", validate=False)
        except Exception as e:
            info.error = f"base64 decode failed: {e}"
            out.append(info)
            continue
        lower = name.lower()
        try:
            if mime == "application/pdf" or lower.endswith(".pdf"):
                info.text, info.method = _pdf_text(data), "pdf"
            elif "wordprocessingml" in mime or lower.endswith(".docx"):
                info.text, info.method = _docx_text(data), "docx"
            elif mime.startswith("text/") or lower.endswith((".txt", ".eml", ".csv")):
                info.text, info.method = data.decode("utf-8", errors="replace"), "text"
            elif mime in IMAGE_MIMES or lower.endswith((".jpg", ".jpeg", ".png", ".gif", ".webp", ".tif", ".tiff", ".bmp")):
                m = mime if mime in IMAGE_MIMES else ("image/png" if lower.endswith(".png") else "image/jpeg")
                try:
                    _vision(data, m, info)
                except llm.LLMUnavailable as e:
                    info.method = "none"
                    info.error = f"image not transcribed ({e}); classified from file name only"
            if info.method == "pdf" and len(info.text.strip()) < 20:
                try:
                    _vision(data, "application/pdf", info)
                except llm.LLMUnavailable as e:
                    info.error = f"PDF has no text layer (scanned) and vision is unavailable ({e})"
        except Exception as e:  # corrupt file
            info.error = f"could not read {name}: {type(e).__name__}: {e}"
        classify(info)
        parse_fields(info)
        out.append(info)
    return out
